#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
交付件验收脚本：检查 docx 交付文件本身
================================================================
为什么要有这个：改完 docx 必须**回读交付文件本身**，
不能只验生成过程 —— 交付路径可能被占用而写失败，或写进去的是旧版。

用法（在项目根目录跑）：
    python _tools\\verify_report_docx.py                       # 默认验「方案报告」
    python _tools\\verify_report_docx.py 分布式抓取协作方案      # 验另一份
    python _tools\\verify_report_docx.py --list                # 看可验目标

每个目标自带一份「关键内容必须在场」的清单（MUST）——
因为不同报告要证明的东西不一样，不能用同一把尺子。
"""
from __future__ import annotations

import io
import re
import sys
import zipfile
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import docx

ROOT = Path(__file__).resolve().parent.parent
REP = ROOT / "_报告"

#: 每个可验目标：md/docx 路径、必须在场的关键内容、引用块探针、最低标题数、最低代码段数
TARGETS: dict[str, dict] = {
    "方案报告": {
        "md": REP / "方案报告.md",
        "docx": REP / "方案报告.docx",
        "must": ["不需要登录", "ajax_getpost.php", "discuz_uid", "500 页", "受理流程",
                 "zoomfile", "archiver", "版权声明", "不易获取信源报告",
                 "CREATE TABLE threads", "reply_sig", "2025-11-11", "两级增量",
                 "pid 集合校验", "token"],
        "quote": "不需要登录",
        "min_heads": 30,
        "min_code": 50,
    },
    "分布式抓取协作方案": {
        "md": REP / "分布式抓取协作方案.md",
        "docx": REP / "分布式抓取协作方案.docx",
        "must": ["不落地", "Crawlab", "authentik", "cloudflared", "FastAPI",
                 "SQLite", "token", "pid 集合", "抽样复核", "熔断", "版本闸",
                 "版权声明", "只读", "租约", "死信", "P1", "Q1", "Huey", "Dramatiq"],
        "quote": "本报告只做设计",
        "min_heads": 25,
        "min_code": 30,
    },
}

FAIL: list[str] = []


def check(cond: bool, msg: str) -> None:
    print(("  [OK]  " if cond else "  [FAIL] ") + msg)
    if not cond:
        FAIL.append(msg)


def verify(name: str) -> int:
    cfg = TARGETS[name]
    MD: Path = cfg["md"]
    DX: Path = cfg["docx"]

    print("=" * 66)
    print(f"交付件验收：{DX.name}")
    print("=" * 66)

    if not DX.exists():
        check(False, f"交付文件不存在：{DX}")
        return 1

    # 1. zip 完整性
    z = zipfile.ZipFile(DX)
    check(z.testzip() is None, "zip 完整性（无损坏条目）")

    # 2. 读回
    d = docx.Document(DX)
    paras = d.paragraphs
    tables = d.tables
    print(f"\n  段落 {len(paras)} 个 / 表格 {len(tables)} 个 / "
          f"表格行合计 {sum(len(t.rows) for t in tables)}")

    alltext = "\n".join(x.text for x in paras)
    for t in tables:
        for row in t.rows:
            for c in row.cells:
                alltext += "\n" + c.text

    def is_code(x) -> bool:
        return bool(x.runs) and x.runs[0].font.name == "Consolas"

    # 3. 正文残留（代码块内的标记属于合法保留，不算残留）
    print("\n--- 正文残留检查（代码块除外）---")
    body_hits = 0
    for x in paras:
        if is_code(x):
            continue
        for m in ("```", "## ", "**", "|---", "\ufffd"):
            if m in x.text:
                print(f"  [FAIL] 残留 {m!r} → {x.text[:90]}")
                body_hits += 1
    check(body_hits == 0, f"正文无 markdown 残留标记（发现 {body_hits} 处）")
    check("\ufffd" not in alltext, "无编码损坏字符 U+FFFD")

    # 4. 关键内容在位
    print("\n--- 关键内容在位 ---")
    for k in cfg["must"]:
        check(k in alltext, f"包含 {k!r}")

    # 5. 标题层级
    heads = [x for x in paras if x.style.name.startswith("Heading")]
    h4 = [x for x in heads if x.style.name == "Heading 4"]
    check(len(heads) >= cfg["min_heads"],
          f"标题数量 {len(heads)} 个（≥{cfg['min_heads']}）")

    # 5b. 标题层级**对齐**：md 用了几级，docx 就得有几级；md 没用过的级别不该凭空出现。
    #     （比"四级标题必须 ≥1"更准：那是转换器的能力代理，
    #       真正要验的是这份文件自己的层级有没有被忠实还原。）
    md_text = MD.read_text(encoding="utf-8")
    md_levels = set()
    for line in md_text.splitlines():
        m = re.match(r"^(#{1,6})\s+\S", line)
        if m:
            md_levels.add(len(m.group(1)))
    docx_levels = {int(x.style.name.split()[-1]) for x in heads
                   if x.style.name.startswith("Heading")
                   and x.style.name.split()[-1].isdigit()}
    for lv in sorted(md_levels):
        n = sum(1 for x in heads if x.style.name == f"Heading {lv}")
        check(n >= 1, f"md 用到 {lv} 级标题 → docx 有 Heading {lv}（{n} 个）")
    for lv in sorted(docx_levels - md_levels):
        check(False, f"docx 出现了 md 没有的 Heading {lv}（层级不忠实）")

    # 6. 代码块与引用块渲染
    code_ps = [x for x in paras if is_code(x)]
    check(len(code_ps) >= cfg["min_code"],
          f"等宽代码段落 {len(code_ps)} 个（≥{cfg['min_code']}）")
    quoted = [x for x in paras if cfg["quote"] in x.text]
    check(len(quoted) >= 1, "引用块内容在位")
    # 引用块的 "> " 前缀必须被剥掉（否则是残留）
    gt = [x for x in paras if x.text.lstrip().startswith("> ")]
    check(len(gt) == 0, f"引用块 '> ' 前缀已剥除（残留 {len(gt)} 处）")

    # 7. 文本量一致（md → docx 不丢内容）
    md = md_text
    md_cn = len(re.findall(r"[\u4e00-\u9fff]", md))
    dx_cn = len(re.findall(r"[\u4e00-\u9fff]", alltext))
    check(md_cn == dx_cn, f"中文字符数一致：md {md_cn} == docx {dx_cn}")

    print("\n" + "=" * 66)
    if FAIL:
        print(f"验收未通过：{len(FAIL)} 项")
        for f in FAIL:
            print("  -", f)
        return 1
    print("验收通过：全部检查项 OK")
    return 0


def main() -> int:
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if "--list" in sys.argv:
        print("可验目标：")
        for k in TARGETS:
            print(f"  - {k}")
        return 0
    if not args:
        name = "方案报告"
    else:
        name = args[0]
        if name not in TARGETS:
            print(f"未知目标 {name!r}；可验目标：{list(TARGETS)}")
            return 2
    return verify(name)


if __name__ == "__main__":
    raise SystemExit(main())
