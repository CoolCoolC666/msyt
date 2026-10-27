#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""md -> docx 转换器（为这份课程报告写的，不是通用 pandoc 替代品）。

为什么不用 pandoc：本机没装（`where pandoc` 空），而这份文档结构很固定，
用 python-docx 直接映射反而更可控，也少一个外部依赖。

处理到的 markdown 子集（够本报告用）：
    # / ## / ###        标题（自动挂中文字体，黑体）
    **加粗**             行内加粗
    `代码`               行内代码（等宽 + 深色）
    | a | b |           表格（首行当表头，Table Grid 边框）
    ---                 分隔线
    1. 2. 3.            有序列表（保留原文编号）
    - x                  无序列表
    *斜体*               行内斜体

中文排版要点（不设就会很难看）：
    python-docx 只设 font.name 时，Word 渲染中文会用默认的"等线"，
    跟正文混排时字号行距全乱。必须同时写 w:eastAsia。

用法：
    py -3 md2docx.py <输入.md> [-o 输出.docx]
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

BODY_CN = "宋体"
HEAD_CN = "黑体"
MONO = "Consolas"
BODY_PT = 11.0

# 行内标记：**粗** `码` *斜*
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\n]+\*)")


def set_cn(run, ascii_font: str, cn_font: str, size: float | None = None,
           bold: bool | None = None, color: RGBColor | None = None) -> None:
    """python-docx 光设 font.name 不管中文，必须补 w:eastAsia。"""
    run.font.name = ascii_font
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn_font)


def add_inline(par, text: str, base_size: float = BODY_PT,
               bold: bool = False) -> None:
    """把一行 markdown 文本拆成若干带格式的 run 塞进段落。

    bold 必须往下传：`**`issue` 类关系`**` 这种加粗里嵌行内代码的写法，
    第一版只剥了外层 **，结果反引号原样留在 Word 正文里（被残留检查抓到）。
    """
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            add_inline(par, piece[2:-2], base_size, bold=True)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            set_cn(r, MONO, BODY_CN, base_size - 0.5, bold=bold,
                   color=RGBColor(0xA3, 0x1D, 0x1D))
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            set_cn(r, "Times New Roman", BODY_CN, base_size, bold=bold)
            r.font.italic = True
        else:
            r = par.add_run(piece)
            set_cn(r, "Times New Roman", BODY_CN, base_size, bold=bold)


def is_table_row(line: str) -> bool:
    return line.lstrip().startswith("|") and line.rstrip().endswith("|")


def split_row(line: str) -> list[str]:
    """★ 转义的竖线（`\\|`）是 markdown 表格里**写竖线的标准做法**（GFM）。

    朴素实现 `split("|")` 会把 `a\\|b` 拆成两个单元格 →
    单元格数超过表头列数时，多出来的那些会被
    「`c >= len(header)` 就跳过」的逻辑**整段丢掉**（内容静默消失）。

    2026-10-27 实测踩到：报告附录 B 有一行写了
    `msyt\\|agentapi\\|example`，docx 里那一行后半段整个不见了
    （md 与 docx 的中文字符数对不上，差 64 个字）。

    → 改成：按**未被转义**的竖线切分，并把 `\\|` 还原成 `|`。
    """
    body = line.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|"):
        body = body[:-1]
    parts = re.split(r"(?<!\\)\|", body)
    return [c.replace("\\|", "|").strip() for c in parts]


def is_sep(line: str) -> bool:
    """markdown 表格的分隔行：|---|---| 或 |:--:|"""
    return is_table_row(line) and set(line.strip().replace("|", "").strip()) <= set("-: ")


def _shade(par, fill: str) -> None:
    """给段落加底纹（代码块用）。"""
    pPr = par._p.get_or_add_pPr()
    shd = pPr.makeelement(qn("w:shd"), {})
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _left_bar(par, color: str) -> None:
    """给段落加左侧竖线（引用块用）。"""
    pPr = par._p.get_or_add_pPr()
    bdr = pPr.makeelement(qn("w:pBdr"), {})
    left = bdr.makeelement(qn("w:left"), {})
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    bdr.append(left)
    pPr.append(bdr)


def convert(md_path: Path, out_path: Path) -> tuple[int, int, int]:
    lines = md_path.read_text(encoding="utf-8").splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.8)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(BODY_PT)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CN)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)

    n_para = n_table = 0
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # ---------- 围栏代码块 ----------
        # 原文（cacx 版）没有这个分支，结果 ```sql 围起来的三行 CREATE TABLE
        # 会被拆成一堆普通段落，Word 里还留着 ``` 标记 —— 表格化的 SQL
        # 方案在 Word 里根本没法看。这里补上：等宽字体 + 浅底纹 + 紧行距。
        if stripped.startswith("```"):
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip())
                i += 1
            i += 1                       # 吃掉收尾的 ```
            for cl in code_lines:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                r = p.add_run(cl if cl.strip() else " ")
                set_cn(r, MONO, BODY_CN, 9, color=RGBColor(0x1A, 0x1A, 0x1A))
                _shade(p, "F4F5F7")
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            n_para += len(code_lines)
            continue

        # ---------- 引用块 ----------
        # 原文也没有。'> ' 开头会被原样当普通段落，Word 里挂着 "> " 很难看。
        # 这里按「左缩进 + 左侧竖线 + 深灰字」处理。
        if stripped.startswith(">"):
            quote_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, " ".join(x for x in quote_lines if x),
                       base_size=BODY_PT - 0.5, bold=False)
            for r_ in p.runs:
                r_.font.color.rgb = RGBColor(0x3A, 0x3A, 0x3A)
            _left_bar(p, "B7C3D6")
            n_para += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 分隔线
        if stripped in ("---", "***", "___"):
            p = doc.add_paragraph()
            set_cn(p.add_run("─" * 34), "Times New Roman", BODY_CN, 9)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # 标题
        # 两处修正（相对 cacx 原版）：
        #  1) 层级从 {1,2,3} 放宽到 {1..6}。原文只认三级，'#### ① 帖级增量'
        #     这种四级小标题会掉进"普通段落"分支，Word 里留下 '#### ' 字样。
        #  2) 标题正文**剥掉行内标记**。原版直接 h.add_run(text)，
        #     于是 '**不做**' 里的星号原样留在标题上（实测 8 处残留）。
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)   # **粗**
            text = re.sub(r"`([^`]+)`", r"\1", text)         # `码`
            text = re.sub(r"\*([^*\n]+)\*", r"\1", text)     # *斜*
            h = doc.add_heading(level=level)
            r = h.add_run(text)
            sizes = {1: 18, 2: 15, 3: 12.5, 4: 11.5, 5: 11, 6: 11}
            set_cn(r, "Times New Roman", HEAD_CN, sizes.get(level, 11), bold=True)
            h.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            h.paragraph_format.space_after = Pt(6)
            n_para += 1
            i += 1
            continue

        # 表格（吃掉表头 + 分隔行 + 若干数据行）
        if is_table_row(stripped):
            header = split_row(stripped)
            j = i + 1
            if j < len(lines) and is_sep(lines[j].strip()):
                j += 1
            rows: list[list[str]] = []
            while j < len(lines) and is_table_row(lines[j].strip()):
                rows.append(split_row(lines[j].strip()))
                j += 1
            t = doc.add_table(rows=1 + len(rows), cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, txt in enumerate(header):
                cell = t.cell(0, c)
                cell.text = ""
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, txt, base_size=10)
                for r_ in p.runs:
                    r_.font.bold = True
            for ri, row in enumerate(rows, start=1):
                for c, txt in enumerate(row):
                    if c >= len(header):
                        continue
                    cell = t.cell(ri, c)
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_inline(p, txt, base_size=10)
            doc.add_paragraph()
            n_table += 1
            i = j
            continue

        # 有序列表：保留原文编号（markdown 编号可能跳号，自己重排会失真）
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.75)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, f"{m.group(1)}. {m.group(2)}")
            n_para += 1
            i += 1
            continue

        # 无序列表
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.75)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, f"· {m.group(1)}")
            n_para += 1
            i += 1
            continue

        # 普通段落：把连续非空行并成一段（markdown 里软换行不该变成新段）
        # 停行条件除了原有的标题/列表/表格/分隔线，还要加上 **引用块** 与
        # **围栏代码块** —— 否则它们会被吸进上一段，Word 里就成了正文的一部分。
        buf = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
                r"^(#{1,6}\s|[-*]\s|\d+\.\s|\||---|>|```)", lines[i].strip()):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(buf))
        n_para += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return n_para, n_table, len(lines)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("-o", "--out", default=None)
    args = ap.parse_args()

    src = Path(args.src)
    out = Path(args.out) if args.out else src.with_suffix(".docx")
    if not src.exists():
        print(f"[FAIL] 找不到 {src}")
        return 1

    n_para, n_table, n_line = convert(src, out)
    print(f"[OK] {src.name} -> {out.name}")
    print(f"     源 {n_line} 行 -> 段落 {n_para} 个、表格 {n_table} 个")
    print(f"     大小 {out.stat().st_size:,} 字节")
    return 0


if __name__ == "__main__":
    sys.exit(main())
